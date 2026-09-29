"""Load submitted documents (.docx, .pdf) as text for the model.

Submissions are free text from different institutions, without a template,
so the loader must not assume any layout. It must also never lose content
silently: a document that loses its author table during loading still yields
a plausible-looking record, and nothing downstream notices. Hence:

- .docx is read from the XML in reading order, including tables, content
  controls, text boxes, footnotes, headers/footers, automatic numbering and
  hyperlink targets. python-docx's ``doc.paragraphs`` skips all of these.
- Superscripts are kept as <sup>…</sup>, because affiliation markers
  ("Rossi²") are otherwise glued to the name ("Rossi2").
- Link targets are kept, because ORCIDs are often only a link behind a name
  or an icon.
- Afterwards every ORCID and DOI found anywhere in the file is checked
  against the loaded text. Anything missing is reported.
- A PDF without a text layer (scanned) is marked fatal instead of sending
  an empty text to the model.
"""

from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Optional

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from validators import DOI_RE, ORCID_RE

SUPPORTED_SUFFIXES = {".docx", ".pdf"}

# Converting these needs LibreOffice or similar; tell the user rather than skip silently.
CONVERT_HINTS = {
    ".doc": "old Word format, please save as .docx",
    ".odt": "OpenDocument, please save as .docx or .pdf",
    ".rtf": "please save as .docx or .pdf",
    ".pages": "please export as .docx or .pdf",
}

MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
HYPERLINK_FIELD_RE = re.compile(r'HYPERLINK\s+"([^"]+)"')

# Below this many characters per page a PDF page is treated as having no text layer.
MIN_CHARS_PER_PDF_PAGE = 20


@dataclass
class LoadedDocument:
    path: Path
    format: str
    text: str
    warnings: list[str] = field(default_factory=list)
    fatal: bool = False  # True: do not send to the model

    def provenance(self) -> dict:
        return {"format": self.format, "loaderWarnings": self.warnings}


def _norm(text: str) -> str:
    return re.sub(r"\s+", "", text).lower()


def _identifiers(text: str) -> set[str]:
    """ORCIDs and DOIs in a text, normalized for comparison."""
    found = {m.group(0).upper() for m in ORCID_RE.finditer(text)}
    for m in DOI_RE.finditer(text):
        found.add(m.group(0).rstrip(".,;:)]}>").lower())
    return found


def markdown_table(rows: list[list[str]]) -> str:
    """Render table rows as Markdown; the first row becomes the header."""
    rows = [[(c or "").replace("|", "\\|").replace("\n", " <br> ").strip() for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + " --- |" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


# --- docx ---------------------------------------------------------------------


def _roman(n: int) -> str:
    out = ""
    for value, sym in ((1000, "M"), (900, "CM"), (500, "D"), (400, "CD"), (100, "C"), (90, "XC"),
                       (50, "L"), (40, "XL"), (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")):
        while n >= value:
            out += sym
            n -= value
    return out


class _Numbering:
    """Renders Word's automatic list numbering, which is not part of the text."""

    FORMATS: dict[str, Callable[[int], str]] = {
        "decimal": str,
        "lowerLetter": lambda n: chr(ord("a") + (n - 1) % 26),
        "upperLetter": lambda n: chr(ord("A") + (n - 1) % 26),
        "lowerRoman": lambda n: _roman(n).lower(),
        "upperRoman": _roman,
    }

    def __init__(self, doc):
        self.levels: dict[tuple[str, int], tuple[str, int, Optional[str]]] = {}
        self.counters: dict[str, dict[int, int]] = {}
        try:
            numbering = doc.part.numbering_part.element
        except (KeyError, NotImplementedError, AttributeError):
            return
        abstract = {}
        for an in numbering.findall(qn("w:abstractNum")):
            lvls = {}
            for lvl in an.findall(qn("w:lvl")):
                fmt = lvl.find(qn("w:numFmt"))
                start = lvl.find(qn("w:start"))
                text = lvl.find(qn("w:lvlText"))
                lvls[int(lvl.get(qn("w:ilvl")))] = (
                    fmt.get(qn("w:val")) if fmt is not None else "decimal",
                    int(start.get(qn("w:val"))) if start is not None else 1,
                    text.get(qn("w:val")) if text is not None else None,
                )
            abstract[an.get(qn("w:abstractNumId"))] = lvls
        for num in numbering.findall(qn("w:num")):
            ref = num.find(qn("w:abstractNumId"))
            if ref is None:
                continue
            for ilvl, spec in abstract.get(ref.get(qn("w:val")), {}).items():
                self.levels[(num.get(qn("w:numId")), ilvl)] = spec

    @staticmethod
    def read_num_pr(num_pr) -> Optional[tuple[str, Optional[int]]]:
        if num_pr is None:
            return None
        num_id_el = num_pr.find(qn("w:numId"))
        ilvl_el = num_pr.find(qn("w:ilvl"))
        num_id = num_id_el.get(qn("w:val")) if num_id_el is not None else None
        ilvl = int(ilvl_el.get(qn("w:val"))) if ilvl_el is not None else None
        return (num_id, ilvl) if num_id is not None or ilvl is not None else None

    def prefix(self, p, style_num: Optional[tuple[str, Optional[int]]] = None) -> str:
        """Numbering set on the paragraph, falling back to its style's numbering."""
        direct = self.read_num_pr(p.find(f"{qn('w:pPr')}/{qn('w:numPr')}"))
        num_id = (direct or (None, None))[0] or (style_num or (None, None))[0]
        ilvl = (direct or (None, None))[1]
        if ilvl is None:
            ilvl = (style_num or (None, 0))[1] or 0
        if num_id is None or num_id == "0":  # none, or explicitly removed
            return ""
        fmt, start, template = self.levels.get((num_id, ilvl), ("bullet", 1, None))
        indent = "  " * ilvl
        if fmt in ("bullet", "none"):
            return f"{indent}- "
        counters = self.counters.setdefault(num_id, {})
        counters[ilvl] = counters.get(ilvl, start - 1) + 1
        for deeper in [k for k in counters if k > ilvl]:
            del counters[deeper]

        # lvlText is Word's label template: "%1." , "(%1)", "%1.%2." ...
        def level_label(match):
            level = int(match.group(1)) - 1
            level_fmt, level_start, _ = self.levels.get((num_id, level), ("decimal", 1, None))
            value = counters.get(level, level_start)
            return self.FORMATS.get(level_fmt, str)(value)

        label = re.sub(r"%(\d)", level_label, template) if template else f"{counters[ilvl]}."
        return f"{indent}{label} "


class _DocxRenderer:
    def __init__(self, doc):
        self.numbering = _Numbering(doc)
        self.heading_levels = {}
        self.style_numbering = {}
        for style in doc.styles.element.findall(qn("w:style")):
            num = _Numbering.read_num_pr(style.find(f"{qn('w:pPr')}/{qn('w:numPr')}"))
            if num:
                self.style_numbering[style.get(qn("w:styleId"))] = num
            name_el = style.find(qn("w:name"))
            name = (name_el.get(qn("w:val")) if name_el is not None else "").lower()
            if name == "title":
                self.heading_levels[style.get(qn("w:styleId"))] = 1
            elif name.startswith("heading "):
                try:
                    self.heading_levels[style.get(qn("w:styleId"))] = int(name.split()[1])
                except ValueError:
                    pass

    # Blocks: paragraphs, tables, content controls
    def blocks(self, container, links: dict) -> str:
        out = []
        for child in container:
            if child.tag == qn("w:p"):
                out.append(self.paragraph(child, links))
            elif child.tag == qn("w:tbl"):
                out.append(self.table(child, links))
            elif child.tag in (qn("w:sdt"), qn("w:customXml")):
                content = child.find(qn("w:sdtContent"))
                out.append(self.blocks(content if content is not None else child, links))
        return "\n".join(b for b in out if b)

    def paragraph(self, p, links: dict) -> str:
        field_urls: list[str] = []
        text = self.inline(p, links, field_urls).strip()
        missing = [u for u in field_urls if _norm(u) not in _norm(text)]
        if missing:
            text += " (" + ", ".join(missing) + ")"
        if not text:
            return ""
        style = p.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
        style_id = style.get(qn("w:val")) if style is not None else None
        level = self.heading_levels.get(style_id)
        prefix = self.numbering.prefix(p, self.style_numbering.get(style_id))
        if level:
            return f"{'#' * min(level, 6)} {prefix}{text}"
        return prefix + text

    def inline(self, el, links: dict, field_urls: list[str]) -> str:
        parts = []
        for child in el:
            tag = child.tag
            if tag == MC_FALLBACK or tag == qn("w:pPr") or tag == qn("w:rPr"):
                continue
            if tag == qn("w:t"):
                parts.append(child.text or "")
            elif tag == qn("w:tab"):
                parts.append("\t")
            elif tag in (qn("w:br"), qn("w:cr")):
                parts.append("\n")
            elif tag == qn("w:noBreakHyphen"):
                parts.append("-")
            elif tag in (qn("w:delText"), qn("w:instrText")):
                if tag == qn("w:instrText"):
                    field_urls.extend(HYPERLINK_FIELD_RE.findall(child.text or ""))
            elif tag == qn("w:footnoteReference"):
                parts.append(f"[^{child.get(qn('w:id'))}]")
            elif tag == qn("w:endnoteReference"):
                parts.append(f"[^e{child.get(qn('w:id'))}]")
            elif tag == qn("w:txbxContent"):
                parts.append("\n" + self.blocks(child, links) + "\n")
            elif tag == qn("w:r"):
                text = self.inline(child, links, field_urls)
                align = child.find(f"{qn('w:rPr')}/{qn('w:vertAlign')}")
                if align is not None and text.strip():
                    val = align.get(qn("w:val"))
                    if val == "superscript":
                        text = f"<sup>{text.strip()}</sup>" + (" " if text.endswith(" ") else "")
                    elif val == "subscript":
                        text = f"<sub>{text.strip()}</sub>" + (" " if text.endswith(" ") else "")
                parts.append(text)
            elif tag == qn("w:hyperlink"):
                text = self.inline(child, links, field_urls)
                url = links.get(child.get(R_ID))
                shown = url[len("mailto:"):] if url and url.startswith("mailto:") else url
                if url and _norm(shown) not in _norm(text) and not url.startswith("#"):
                    text = f"{text} ({url})"
                parts.append(text)
            elif tag == qn("w:fldSimple"):
                field_urls.extend(HYPERLINK_FIELD_RE.findall(child.get(qn("w:instr")) or ""))
                parts.append(self.inline(child, links, field_urls))
            else:
                # smartTag, ins, sdt, AlternateContent/Choice, drawing, ...
                parts.append(self.inline(child, links, field_urls))
        return "".join(parts)

    def table(self, tbl, links: dict) -> str:
        return markdown_table([
            [self.blocks(tc, links) for tc in tr.findall(qn("w:tc"))]
            for tr in tbl.findall(qn("w:tr"))
        ])


def _external_links(part) -> dict:
    links = {}
    for r_id, rel in part.rels.items():
        if rel.is_external:
            links[r_id] = rel.target_ref
    return links


def _related_parts(doc, suffix: str) -> list:
    return [rel.target_part for rel in doc.part.rels.values()
            if not rel.is_external and rel.reltype.endswith(suffix)]


def _notes(renderer: _DocxRenderer, part, note_tag: str, marker: str) -> list[str]:
    root = etree.fromstring(part.blob)
    links = _external_links(part)
    out = []
    for note in root.findall(qn(note_tag)):
        if note.get(qn("w:type")) in ("separator", "continuationSeparator", "continuationNotice"):
            continue
        text = renderer.blocks(note, links).strip()
        if text:
            out.append(f"[^{marker}{note.get(qn('w:id'))}]: {text}")
    return out


def _load_docx(path: Path) -> LoadedDocument:
    doc = Document(str(path))
    renderer = _DocxRenderer(doc)
    sections = [renderer.blocks(doc.element.body, _external_links(doc.part))]

    notes = []
    for part in _related_parts(doc, "/footnotes"):
        notes += _notes(renderer, part, "w:footnote", "")
    for part in _related_parts(doc, "/endnotes"):
        notes += _notes(renderer, part, "w:endnote", "e")
    if notes:
        sections.append("## Footnotes\n" + "\n".join(notes))

    # Letterheads often carry the institution. Headers repeat per section,
    # so keep each distinct text once.
    seen, margins = set(), []
    for suffix in ("/header", "/footer"):
        for part in _related_parts(doc, suffix):
            text = renderer.blocks(etree.fromstring(part.blob), _external_links(part)).strip()
            if text and text not in seen:
                seen.add(text)
                margins.append(text)
    if margins:
        sections.append("## Page header/footer\n" + "\n".join(margins))

    text = "\n\n".join(s for s in sections if s.strip())
    loaded = LoadedDocument(path=path, format="docx", text=text)
    _check_docx_completeness(path, loaded)
    return loaded


def _check_docx_completeness(path: Path, loaded: LoadedDocument) -> None:
    """Compare identifiers in the raw file with the loaded text.

    Deliberately independent of the renderer above: every text node of every
    XML part plus every external link target, so a renderer bug shows up
    here instead of as a silently thinner record.
    """
    raw_chunks = []
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if not name.startswith("word/"):
                continue
            data = z.read(name)
            if name.endswith(".rels"):
                raw_chunks += re.findall(r'Target="([^"]+)"', data.decode("utf-8", "replace"))
            elif name.endswith(".xml"):
                try:
                    root = etree.fromstring(data)
                except etree.XMLSyntaxError:
                    continue
                for p in root.iter(qn("w:p")):
                    raw_chunks.append("".join(t.text or "" for t in p.iter(qn("w:t"), qn("w:instrText"))))
    raw_ids = _identifiers("\n".join(raw_chunks))
    loaded_text = _norm(re.sub(r"</?su[bp]>", "", loaded.text))
    missing = sorted(i for i in raw_ids if _norm(i) not in loaded_text)
    if missing:
        loaded.warnings.append(
            "identifiers in the file but not in the loaded text (content lost while loading): "
            + ", ".join(missing)
        )


# --- pdf ----------------------------------------------------------------------


def _load_pdf(path: Path) -> LoadedDocument:
    # Imported lazily: pdf_loader imports this module's helpers.
    from pdf_loader import load_pdf
    return load_pdf(path)


# --- entry points ---------------------------------------------------------------


LOADERS = {".docx": _load_docx, ".pdf": _load_pdf}


def is_input_file(path: Path) -> bool:
    """Files that can be submissions: supported formats and those we can give a
    conversion hint for. Skips Word lock files (~$x.docx), hidden files,
    directories and everything else (e.g. finished records as .json)."""
    return (path.is_file()
            and not path.name.startswith(("~$", "."))
            and path.suffix.lower() in SUPPORTED_SUFFIXES | CONVERT_HINTS.keys())


def load_document(path: Path) -> LoadedDocument:
    loader = LOADERS.get(path.suffix.lower())
    if loader is None:
        hint = CONVERT_HINTS.get(path.suffix.lower(), "unsupported format")
        return LoadedDocument(path=path, format=path.suffix.lower().lstrip("."), text="",
                              warnings=[hint], fatal=True)
    loaded = loader(path)
    if not loaded.text.strip() and not loaded.fatal:
        loaded.fatal = True
        loaded.warnings.append("no text found in document")
    return loaded


if __name__ == "__main__":
    # Preview what the model will see, without calling it:
    #   uv run python documents.py data/antrag.pdf [more files or folders ...]
    import sys

    if len(sys.argv) < 2:
        sys.exit("usage: documents.py FILE_OR_FOLDER [...]")
    targets = []
    for arg in sys.argv[1:]:
        path = Path(arg)
        targets += sorted(p for p in path.iterdir() if is_input_file(p)) if path.is_dir() else [path]
    for path in targets:
        doc = load_document(path)
        status = "[SKIP] would be skipped" if doc.fatal else ("[WARN] with warnings" if doc.warnings else "[OK]")
        print(f"\n{'═' * 80}\n{path.name}  ({doc.format}, {len(doc.text)} chars)  {status}")
        for warning in doc.warnings:
            print(f"  [WARN] {warning}")
        print("─" * 80)
        print(doc.text or "(no text)")

